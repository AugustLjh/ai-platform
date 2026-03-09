"""
File Parsers for Knowledge Base
"""
import csv
import json
import os
from abc import ABC, abstractmethod
from typing import Optional
from io import BytesIO, StringIO
from xml.etree import ElementTree as ET


class FileParser(ABC):
    """文件解析器抽象基类"""

    @abstractmethod
    async def parse(self, file_content: bytes, filename: str) -> str:
        """
        解析文件内容为文本

        Args:
            file_content: 文件二进制内容
            filename: 文件名

        Returns:
            解析后的文本内容
        """
        pass

    @staticmethod
    def get_parser(filename: str) -> Optional["FileParser"]:
        """
        根据文件名获取对应的解析器

        Args:
            filename: 文件名

        Returns:
            对应的解析器，如果不支持则返回None
        """
        ext = os.path.splitext(filename)[1].lower()

        text_parser = TextParser()
        markdown_parser = MarkdownParser()
        parsers = {
            ".txt": text_parser,
            ".text": text_parser,
            ".log": text_parser,
            ".md": markdown_parser,
            ".markdown": markdown_parser,
            ".pdf": PDFParser(),
            ".html": HTMLParser(),
            ".htm": HTMLParser(),
            ".csv": CSVParser(),
            ".tsv": CSVParser(delimiter="\t"),
            ".json": JSONParser(),
            ".jsonl": JSONLParser(),
            ".yaml": YAMLParser(),
            ".yml": YAMLParser(),
            ".xml": XMLParser(),
            ".rtf": RTFParser(),
            ".docx": DocxParser(),
            ".pptx": PptxParser(),
            ".xlsx": XlsxParser(),
        }

        return parsers.get(ext)

    @staticmethod
    def is_supported(filename: str) -> bool:
        """检查文件类型是否支持"""
        return FileParser.get_parser(filename) is not None


def _decode_bytes(file_content: bytes) -> str:
    encodings = ["utf-8", "gbk", "gb2312", "latin-1"]
    for encoding in encodings:
        try:
            return file_content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return file_content.decode("utf-8", errors="ignore")


class TextParser(FileParser):
    """纯文本解析器"""

    async def parse(self, file_content: bytes, filename: str) -> str:
        """解析文本文件"""
        return _decode_bytes(file_content)


class MarkdownParser(FileParser):
    """Markdown解析器"""

    async def parse(self, file_content: bytes, filename: str) -> str:
        """解析Markdown文件（保留原始格式）"""
        return _decode_bytes(file_content)


class CSVParser(FileParser):
    """CSV/TSV解析器"""

    def __init__(self, delimiter: Optional[str] = None):
        self.delimiter = delimiter

    async def parse(self, file_content: bytes, filename: str) -> str:
        text = _decode_bytes(file_content)
        if not text.strip():
            return ""

        try:
            if self.delimiter:
                reader = csv.reader(StringIO(text), delimiter=self.delimiter)
            else:
                dialect = csv.Sniffer().sniff(text[:4096])
                reader = csv.reader(StringIO(text), dialect)
        except csv.Error:
            reader = csv.reader(StringIO(text))
        lines = []
        for row in reader:
            if not row:
                continue
            lines.append("\t".join(str(cell) for cell in row))

        return "\n".join(lines) if lines else text


class JSONParser(FileParser):
    """JSON解析器"""

    async def parse(self, file_content: bytes, filename: str) -> str:
        text = _decode_bytes(file_content)
        if not text.strip():
            return ""
        try:
            data = json.loads(text)
            return json.dumps(data, ensure_ascii=False, indent=2)
        except (json.JSONDecodeError, ValueError):
            return text


class JSONLParser(FileParser):
    """JSONL解析器"""

    async def parse(self, file_content: bytes, filename: str) -> str:
        text = _decode_bytes(file_content)
        if not text.strip():
            return ""
        lines = []
        for line in text.splitlines():
            stripped = line.strip()
            if not stripped:
                continue
            try:
                data = json.loads(stripped)
                lines.append(json.dumps(data, ensure_ascii=False))
            except (json.JSONDecodeError, ValueError):
                lines.append(stripped)
        return "\n".join(lines)


class YAMLParser(FileParser):
    """YAML解析器（作为纯文本处理）"""

    async def parse(self, file_content: bytes, filename: str) -> str:
        return _decode_bytes(file_content)


class XMLParser(FileParser):
    """XML解析器"""

    async def parse(self, file_content: bytes, filename: str) -> str:
        text = _decode_bytes(file_content)
        if not text.strip():
            return ""
        try:
            root = ET.fromstring(text)
            content = "\n".join(t.strip() for t in root.itertext() if t.strip())
            return content or text
        except ET.ParseError:
            return text


class RTFParser(FileParser):
    """RTF解析器"""

    async def parse(self, file_content: bytes, filename: str) -> str:
        try:
            from striprtf.striprtf import rtf_to_text
        except ImportError:
            raise ImportError(
                "striprtf not installed. Install it with: pip install striprtf"
            )
        text = _decode_bytes(file_content)
        return rtf_to_text(text).strip()


class PDFParser(FileParser):
    """PDF解析器"""

    async def parse(self, file_content: bytes, filename: str) -> str:
        """
        解析PDF文件

        需要安装：pip install PyPDF2
        """
        try:
            from PyPDF2 import PdfReader
            from io import BytesIO
        except ImportError:
            raise ImportError(
                "PyPDF2 not installed. Install it with: pip install PyPDF2"
            )

        try:
            pdf_file = BytesIO(file_content)
            reader = PdfReader(pdf_file)

            text_parts = []
            for page_num, page in enumerate(reader.pages, 1):
                text = page.extract_text()
                if text.strip():
                    text_parts.append(f"=== Page {page_num} ===\n{text}")

            return "\n\n".join(text_parts)

        except Exception as e:
            raise ValueError(f"Failed to parse PDF: {str(e)}")


class DocxParser(FileParser):
    """DOCX解析器"""

    async def parse(self, file_content: bytes, filename: str) -> str:
        try:
            from docx import Document as DocxDocument
        except ImportError:
            raise ImportError(
                "python-docx not installed. Install it with: pip install python-docx"
            )

        try:
            doc = DocxDocument(BytesIO(file_content))
            parts = []

            for para in doc.paragraphs:
                text = para.text.strip()
                if text:
                    parts.append(text)

            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        parts.append("\t".join(cells))

            return "\n".join(parts)
        except Exception as e:
            raise ValueError(f"Failed to parse DOCX: {str(e)}")


class PptxParser(FileParser):
    """PPTX解析器"""

    async def parse(self, file_content: bytes, filename: str) -> str:
        try:
            from pptx import Presentation
        except ImportError:
            raise ImportError(
                "python-pptx not installed. Install it with: pip install python-pptx"
            )

        try:
            prs = Presentation(BytesIO(file_content))
            parts = []
            for idx, slide in enumerate(prs.slides, 1):
                slide_texts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        slide_texts.append(shape.text.strip())
                if slide_texts:
                    parts.append(f"=== Slide {idx} ===")
                    parts.append("\n".join(slide_texts))
            return "\n\n".join(parts)
        except Exception as e:
            raise ValueError(f"Failed to parse PPTX: {str(e)}")


class XlsxParser(FileParser):
    """XLSX解析器"""

    async def parse(self, file_content: bytes, filename: str) -> str:
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ImportError(
                "openpyxl not installed. Install it with: pip install openpyxl"
            )

        try:
            workbook = load_workbook(BytesIO(file_content), data_only=True, read_only=True)
            parts = []
            for sheet in workbook.worksheets:
                parts.append(f"=== Sheet: {sheet.title} ===")
                for row in sheet.iter_rows(values_only=True):
                    if row is None:
                        continue
                    cells = ["" if v is None else str(v) for v in row]
                    if any(cell.strip() for cell in cells):
                        parts.append("\t".join(cells))
            return "\n".join(parts)
        except Exception as e:
            raise ValueError(f"Failed to parse XLSX: {str(e)}")


class HTMLParser(FileParser):
    """HTML解析器"""

    async def parse(self, file_content: bytes, filename: str) -> str:
        """
        解析HTML文件，提取纯文本

        需要安装：pip install beautifulsoup4
        """
        try:
            from bs4 import BeautifulSoup
        except ImportError:
            raise ImportError(
                "beautifulsoup4 not installed. Install it with: pip install beautifulsoup4"
            )

        try:
            # 尝试检测编码
            encodings = ['utf-8', 'gbk', 'gb2312', 'latin-1']
            html_content = None

            for encoding in encodings:
                try:
                    html_content = file_content.decode(encoding)
                    break
                except UnicodeDecodeError:
                    continue

            if html_content is None:
                html_content = file_content.decode('utf-8', errors='ignore')

            # 使用BeautifulSoup提取文本
            soup = BeautifulSoup(html_content, 'html.parser')

            # 移除script和style标签
            for script in soup(['script', 'style']):
                script.decompose()

            # 获取文本
            text = soup.get_text(separator='\n', strip=True)

            # 清理多余的空行
            lines = [line.strip() for line in text.split('\n') if line.strip()]
            return '\n'.join(lines)

        except Exception as e:
            raise ValueError(f"Failed to parse HTML: {str(e)}")
